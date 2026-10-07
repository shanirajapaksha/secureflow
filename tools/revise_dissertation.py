from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\HP\Downloads\ai_based_nids_final_dissertation.docx")
OUTPUT = ROOT / "ai_based_nids_final_dissertation_revised.docx"
ASSETS = ROOT / "artifacts"


def set_text(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def replace_start(document, prefix, text):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_text(paragraph, text)
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def find_exact(document, text):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def replace_with_picture(paragraph, image_path, width=6.3):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))


def insert_before(reference, text="", style=None):
    new_p = OxmlElement("w:p")
    reference._p.addprevious(new_p)
    paragraph = reference._parent.add_paragraph()
    paragraph._p.getparent().remove(paragraph._p)
    new_p.addnext(paragraph._p)
    if style:
        paragraph.style = style
    if text:
        paragraph.add_run(text)
    return paragraph


def insert_picture_before(reference, image_path, caption, width=6.3):
    picture = insert_before(reference)
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(image_path), width=Inches(width))
    cap = insert_before(reference, caption, "Body Text")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def fill_table(table, rows):
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = str(value)


doc = Document(SOURCE)

# Abstract and front matter: replace unsupported model, alert, and storage claims.
replace_start(
    doc,
    "This research presents the design, implementation",
    "This research presents the design, implementation, and evaluation of a hybrid AI-based Network Intrusion Detection System (NIDS). The machine-learning pipeline was developed using CICIDS2017 flow data, while live packet captures are collected through TShark/Wireshark-compatible PCAP capture and converted into flow records using CICFlowMeter. The implemented preprocessing workflow performs stratified train/test splitting, training-only missing-value treatment, Isolation Forest outlier filtering, correlation-based feature reduction, StandardScaler normalization, and capped multiclass oversampling.",
)
replace_start(
    doc,
    "Four supervised machine learning classifiers",
    "Three supervised classifiers were evaluated: Decision Tree, Random Forest, and Logistic Regression. On the 504,402-record holdout test set, Decision Tree achieved the highest weighted F1-score (0.997958) and accuracy (0.997960), and was therefore serialized as the deployed model. Random Forest achieved 0.997760 accuracy and 0.997714 weighted F1, while Logistic Regression achieved 0.963474 accuracy and 0.976133 weighted F1. The final model performs multiclass classification across BENIGN and fourteen CICIDS2017 attack categories.",
)
replace_start(
    doc,
    "To translate model intelligence into actionable operational workflows",
    "The operational solution combines a FastAPI backend with a React 18 and TypeScript single-page dashboard. PostgreSQL is used as the persistent data store for user accounts, traffic and prediction records, alerts, blocking actions, and history/audit information. The system supports login and signup, dashboard alert delivery, automatic blocking based on configured detection policy, manual administrator blocking, and review of historical events. Live and uploaded CICFlowMeter-compatible CSV data are processed through the same hybrid signature and machine-learning detection service.",
)
replace_start(
    doc,
    "Keywords:",
    "Keywords: Network Intrusion Detection System (NIDS), Artificial Intelligence, CICIDS2017, Decision Tree, Random Forest, Hybrid Detection, PostgreSQL, Real-Time Monitoring, Automated Blocking.",
)

# Front lists and appendix labels.
for paragraph in doc.paragraphs:
    if paragraph.text.strip() == "Figure 6.1: Confusion Matrix for Random Forest Classifier on Evaluation Dataset":
        set_text(paragraph, "Figure 6.1: Binary Confusion Matrix for the Final Saved Classifier")
    elif paragraph.text.strip() == "Appendix B: Core Machine Learning Training Script (train_random_forest.py)":
        set_text(paragraph, "Appendix B: Model Training Notebook (04.train_model.ipynb)")

# Chapter 3 methodology.
replace_start(
    doc,
    "+-----------------------------------------------------------------------+\n|                      PHASE 1:",
    "PHASE 1 – Requirements and feasibility: stakeholder requirements, survey evidence, technical feasibility, and ethical constraints.\n\nPHASE 2 – Data engineering: CICIDS2017 consolidation and cleaning, 80/20 stratified split, training-only imputation and outlier removal, correlation pruning, scaling, and capped multiclass oversampling.\n\nPHASE 3 – Model engineering: comparative training of Decision Tree, Random Forest, and Logistic Regression, followed by multiclass evaluation and Joblib serialization of the highest-F1 model.\n\nPHASE 4 – System integration: FastAPI prediction services, PostgreSQL persistence, React dashboard, live TShark/CICFlowMeter ingestion, alerts, automatic/manual blocking, authentication, and history views.\n\nPHASE 5 – Verification: holdout evaluation, API unit tests, live-capture validation, dashboard route checks, and inference timing.",
)
replace_start(doc, "Software Tools: Python 3.10", "Software Tools: The verified development environment uses Python 3.14.3, while project documentation specifies Python 3.9 or later. The implementation uses FastAPI, Scikit-learn, Pandas, NumPy, Joblib, React 18, TypeScript, Vite, Tailwind CSS, Recharts, TShark, CICFlowMeter, and PostgreSQL.")
replace_start(doc, "Machine Learning Classifier Engineering:", "Machine Learning Classifier Engineering: Implement and compare Decision Tree, Random Forest, and Logistic Regression under identical preprocessing and holdout-evaluation conditions, then serialize the highest weighted-F1 model for deployment.")
replace_start(doc, "An intuitive dashboard interface displaying", "An intuitive dashboard interface displaying color-coded severity information, persistent alert history, and administrator controls for reviewing and manually blocking detected sources. Alert generation and automatic blocking are also supported according to configured detection policies.")
replace_start(
    doc,
    "To establish empirical domain requirements",
    "Two structured questionnaires were used to gather requirements from general users and technical stakeholders. The dissertation records 52 general respondents and 32 technical respondents. However, the repository does not contain the Google Forms configuration or distribution records; therefore, the exact sharing channels, anonymity setting, and consent wording should be confirmed from the original survey owners before final submission.",
)
replace_start(doc, "Preferred Alert Notification Channels:", "Preferred Alert Notification Channels: The recorded survey summary identifies the centralized dashboard as the leading preference, followed by email and SMS. In the implemented system, alerts are persisted in PostgreSQL and presented through the dashboard/history workflow; any external email or SMS channel should be described only if its configured service can be demonstrated.")

pipeline = replace_start(
    doc,
    "+---------------------------------------+\n                      |       Raw CICIDS2017 Dataset",
    "Verified preprocessing pipeline figure.",
)
replace_with_picture(pipeline, ASSETS / "dissertation_assets" / "preprocessing_pipeline.png")
replace_start(doc, "As illustrated in Figure 3.3", "As illustrated in Figure 3.3, the cleaned CICIDS2017 working dataset contains 2,522,009 records and 70 input features. It contains 2,096,134 BENIGN records and 425,875 attack records distributed across fourteen attack categories. The workflow prevents evaluation leakage by performing the train/test split before fitting imputation, outlier filtering, correlation removal, scaling, and oversampling operations.")
replace_start(doc, "Data Ingestion & Consolidation:", "Data Ingestion and Cleaning: Eight CICIDS2017 CSV files were consolidated and cleaned to produce 2,522,009 usable records. Column names were normalized and the target Label field was retained as a multiclass label.")
replace_start(doc, "Missing & Infinite Value Handling:", "Missing and Infinite Value Handling: After the stratified split, positive and negative infinity values were converted to missing values. Means learned exclusively from the training partition were used to fill missing values in both training and test data.")
replace_start(doc, "Duplicate Record Elimination:", "Train/Test Split: A stratified 80/20 split produced 2,017,607 training records and 504,402 independent test records before train-only transformations.")
replace_start(doc, "Low-Variance Feature Pruning:", "Feature Pruning: No zero-variance columns were found at this stage. Correlation analysis fitted on the training partition removed 23 highly correlated columns, reducing the feature vector from 70 to 47 dimensions.")
replace_start(doc, "Isolation Forest Outlier Filtering:", "Isolation Forest Outlier Filtering: Isolation Forest was applied only to the training partition. The training set decreased from 2,017,607 to 1,977,254 records. Immediately before oversampling it contained 1,643,718 BENIGN and 333,536 attack records.")
replace_start(doc, "where $E(h(x))", "Outlier detection was deliberately excluded from the test partition so that evaluation remained representative of unseen traffic.")
replace_start(doc, "Mathematical Feature Scaling & Class Rebalancing:", "Feature Scaling and Class Rebalancing: StandardScaler was fitted on the filtered training data and then applied to both partitions. Multiclass oversampling was performed only on training data, with minority classes capped at 100,000 samples to control memory use.")
replace_start(doc, "Standardization ensures that", "This order ensures that test-set statistics do not influence training-time transformations and that the holdout set remains unmodified for evaluation.")
replace_start(doc, "To resolve severe class imbalance", "After capped oversampling, the final training matrix contained 2,979,959 rows and 47 float32 features. The test matrix remained unchanged at 504,402 rows. Labels were encoded into fifteen categories: BENIGN and fourteen attack types.")

# Chapter 4 architecture and persistence.
replace_start(doc, "The system design phase translates", "The system design translates requirements into an implemented multi-tier architecture. This chapter documents capture and upload inputs, the FastAPI detection service, PostgreSQL persistence, authentication, alerts, blocking controls, history, and the React presentation layer.")
arch = replace_start(
    doc,
    "+-----------------------------------------------------------------------+\n    |                    DATA COLLECTION",
    "DATA SOURCES\n• Uploaded CICFlowMeter-compatible CSV files\n• Live TShark PCAP/PCAPNG captures converted by CICFlowMeter\n\nFASTAPI APPLICATION\n• Input validation and column normalization\n• Saved preprocessing artifacts and multiclass Decision Tree inference\n• Signature-rule checks and severity assignment\n• Authentication, alerts, blocking, and history services\n\nPOSTGRESQL\n• Users and authentication metadata\n• Traffic flows and prediction results\n• Alerts, automatic/manual block actions, and audit/history records\n\nREACT 18 / TYPESCRIPT DASHBOARD\n• Dashboard, Network Traffic, Real-Time Capture, Threat Alerts, Settings, login/signup, blocking controls, and history views",
)
replace_start(doc, "Data Collection Tier:", "Data Collection Tier: Accepts uploaded CICIDS2017/CICFlowMeter-compatible CSV files and live PCAP/PCAPNG traffic captured with TShark. CICFlowMeter converts packet captures into the flow-level features expected by the model.")
replace_start(doc, "AI Engine Tier:", "Application and Detection Tier: FastAPI validates feature coverage, applies saved preprocessing artifacts, executes signature rules, and performs multiclass inference using the deployed DecisionTreeClassifier. Detection outcomes can generate alerts and automatic blocking actions according to configured policy.")
replace_start(doc, "Presentation Tier:", "Presentation and Persistence Tier: React 18 and TypeScript provide operational views, while PostgreSQL stores user accounts, traffic and prediction data, alerts, block lists/actions, and historical/audit records. Administrators can review alerts and perform manual blocking.")
replace_start(doc, "Administrator Workflow:", "Administrator Workflow: Users register or sign in, view the dashboard and historical records, inspect alert details, and manually block suspicious sources. The system can also invoke an automatic block action when a detection satisfies the configured severity or policy threshold. Model retraining remains a separate notebook workflow unless a dedicated retraining service is deployed.")
replace_start(
    doc,
    "+-----------------------------------------------------------------+\n       |                        AI-BASED NIDS",
    "ACTORS AND USE CASES\n\nNetwork Traffic Source\n• Upload CICFlowMeter-compatible CSV\n• Supply live TShark capture for flow conversion\n\nRegistered User / Administrator\n• Sign up and log in\n• View dashboard, network traffic, alerts, and history\n• Inspect detection details\n• Manually block or unblock a source\n\nAutomated System\n• Validate and preprocess traffic\n• Apply signature rules and ML classification\n• Persist results and alerts in PostgreSQL\n• Apply configured automatic-block policy\n• Record all security and administrative actions in history",
)
replace_start(doc, "The Sequence Diagram illustrates", "The sequence begins when an authenticated user uploads a flow CSV or starts a live capture. FastAPI validates and preprocesses the data, applies signature rules and the saved ML model, stores traffic, prediction, and alert records in PostgreSQL, evaluates the automatic-block policy, and returns JSON to the dashboard. Administrators may subsequently issue a manual block request, and each action is appended to the history/audit trail.")
replace_start(
    doc,
    "Network       Network      NIDS Core",
    "SEQUENCE OF OPERATIONS\n1. User signs up or logs in; the authentication service verifies PostgreSQL account data.\n2. User uploads a CSV or starts a TShark capture.\n3. CICFlowMeter converts a PCAP/PCAPNG file into compatible flow features when live capture is used.\n4. FastAPI validates, preprocesses, and sends the feature matrix to signature and ML detection.\n5. Traffic, predictions, and any generated alerts are stored in PostgreSQL.\n6. The policy engine applies an automatic block when its configured conditions are satisfied.\n7. The API returns the result to the React dashboard.\n8. An administrator may request a manual block/unblock action.\n9. Alerts and administrative actions are appended to the persistent audit/history record.",
)
replace_start(doc, "The Activity Diagram tracks", "The activity flow covers authentication, traffic acquisition, CICFlowMeter feature extraction, validation, preprocessing, signature and ML detection, PostgreSQL persistence, alert generation, automatic policy-based blocking, optional manual administrator blocking, dashboard presentation, and history review.")
replace_start(doc, "The database schema models entities", "PostgreSQL is the implemented persistence layer. The logical schema models user accounts, authentication/roles, captured or uploaded traffic, model predictions, alerts, automatic and manual blocking actions, model metadata, and immutable audit/history entries.")
replace_start(doc, "The database structure was normalized", "The PostgreSQL schema is normalized to reduce duplication and preserve referential integrity. User, traffic, prediction, alert, block-action, and audit-history entities are stored independently and joined through primary and foreign keys.")
replace_start(doc, "First Normal Form (1NF):", "First Normal Form (1NF): Each column stores one atomic value; repeating alerts, roles, or blocking actions are stored as separate rows.")
replace_start(doc, "Second Normal Form (2NF):", "Second Normal Form (2NF): Non-key fields depend on the complete primary key, and relationship tables are used where many-to-many associations are required.")
replace_start(doc, "Third Normal Form (3NF):", "Third Normal Form (3NF): User, traffic, prediction, alert, block, and history attributes are separated so that non-key fields depend only on their entity key.")
schema = replace_start(
    doc,
    "ACTOR (\n",
    "POSTGRESQL LOGICAL SCHEMA\n\nUSERS(user_id PK, name, email UNIQUE, password_hash, role, created_at)\nTRAFFIC_FLOWS(flow_id PK, captured_at, source_ip, destination_ip, protocol, feature_payload, capture_source)\nPREDICTIONS(prediction_id PK, flow_id FK, model_id FK, predicted_type, traffic_status, detection_method, confidence, created_at)\nALERTS(alert_id PK, prediction_id FK, severity, reason, status, created_at)\nBLOCK_ACTIONS(block_id PK, alert_id FK, source_ip, mode, requested_by FK, status, created_at)\nMODEL_METADATA(model_id PK, model_name, artifact_path, trained_at, metrics)\nAUDIT_HISTORY(history_id PK, user_id FK, action_type, entity_type, entity_id, details, created_at)\n\nThe mode field distinguishes AUTO from MANUAL blocking. Foreign keys connect every operational event to its source flow, prediction, alert, administrator, and model version.",
)
replace_start(doc, "Figure 4.7:", "Figure 4.7: PostgreSQL Logical Schema Specification")

# Functional requirements table: reflect confirmed features.
fill_table(
    doc.tables[3],
    [
        ["ID", "Requirement Name", "Detailed Description", "Priority"],
        ["FR-01", "Traffic Ingestion", "Accept uploaded flow CSVs and live TShark/CICFlowMeter captures.", "High"],
        ["FR-02", "Hybrid Detection", "Apply signature rules and the serialized multiclass ML model.", "High"],
        ["FR-03", "PostgreSQL Persistence", "Persist users, traffic, predictions, alerts, blocks, and history.", "High"],
        ["FR-04", "Alerts", "Create and display alerts with attack type, method, reason, and severity.", "High"],
        ["FR-05", "Automatic/Manual Blocking", "Apply policy-based automatic blocks and administrator-requested manual blocks.", "High"],
        ["FR-06", "Authentication", "Provide login, signup, authorization, and account records.", "High"],
        ["FR-07", "History and Audit", "Allow authorized users to review stored detections and administrative actions.", "High"],
    ],
)
doc.tables[2].cell(3, 2).text = (
    "Persist alerts in PostgreSQL and present them through the dashboard and history views; "
    "external email/SMS delivery should be documented only for configured channels."
)

# Chapter 5 implementation details.
replace_start(doc, "AI-Based NIDS Core (Backend Engine):", "AI-Based NIDS Core: Manages model and artifact loading, CSV validation, preprocessing, hybrid detection, live TShark/CICFlowMeter processing, PostgreSQL writes, alert generation, authentication, blocking operations, and history/audit services.")
replace_start(doc, "AI-Based NIDS Dashboard (Frontend UI):", "AI-Based NIDS Dashboard: A React 18 and TypeScript SPA that consumes FastAPI JSON endpoints and provides dashboard, network traffic, real-time capture, threat alert, settings, authentication, block-management, and history interfaces.")
replace_start(doc, "Python Runtime:", "Python Runtime: Python 3.14.3 was detected in the verified environment; the project documentation supports Python 3.9 or later through an isolated virtual environment.")
replace_start(doc, "Core ML Libraries:", "Core Backend and ML Libraries: FastAPI 0.135.1, Pandas 3.0.1, NumPy 2.4.3, Scikit-learn 1.8.0, SciPy 1.17.1, Joblib 1.5.3, Uvicorn, and PostgreSQL connectivity used by the persistence layer.")
replace_start(doc, "Frontend Runtime:", "Frontend Runtime: React 18.3, TypeScript 5.8, Vite 5.4, Tailwind CSS 3.4, Recharts 2.15, React Router, and Shadcn/Radix UI components.")
replace_start(doc, "The system supports decoupled deployment", "The system uses decoupled deployment: the React SPA is served independently, FastAPI exposes RESTful JSON endpoints, PostgreSQL provides durable persistence, and TShark/CICFlowMeter support live capture conversion. This separation allows the dashboard, API, detection engine, database, and capture tools to be deployed and maintained independently.")
replace_start(
    doc,
    "+-----------------------------------------------------------------+\n |                           USER BROWSER",
    "USER BROWSER\nReact 18 / TypeScript SPA: authentication, dashboard, alerts, blocking, and history\n\nHTTP / JSON REST API\n\nFASTAPI APPLICATION\nValidation, preprocessing, signature rules, ML inference, alert and blocking services\n\nPOSTGRESQL DATABASE\nPersistent users, flows, predictions, alerts, automatic/manual block actions, and history\n\nCAPTURE TOOLCHAIN\nTShark packet capture → PCAP/PCAPNG → CICFlowMeter → flow CSV → detection API",
)
replace_start(doc, "The initial preprocessing phase merges", "The data pipeline is implemented as a series of Jupyter notebooks: dataset acquisition, cleaning, leakage-safe preprocessing, model comparison, intrusion prediction, final evaluation, and signature-rule analysis.")
replace_start(doc, "Numerical features are standardized", "The preprocessing notebook performs a stratified split before fitting train-only transformations. Isolation Forest filtering and oversampling are applied only to training data; StandardScaler and selected feature metadata are serialized for consistent API inference.")
replace_start(doc, "The Random Forest ensemble classifier is implemented", "The training notebook compares Decision Tree, Random Forest, and Logistic Regression using accuracy, weighted precision, weighted recall, weighted F1, and training time. Decision Tree achieved the highest weighted F1 and was serialized with Joblib as best_nids_model.joblib; model_info.joblib stores the comparison table.")

# Replace fabricated code excerpts with accurate concise excerpts.
for p in doc.paragraphs:
    text = p.text.strip()
    if text.startswith("import os\nimport pandas as pd"):
        set_text(p, "# See notebooks/02.clean_dataset.ipynb\n# Consolidate the eight CICIDS2017 CSV files, normalize labels and columns,\n# remove invalid/duplicate rows, and export cicids2017_cleaned_nids.csv.")
    elif text.startswith("import pandas as pd\nimport numpy as np"):
        set_text(p, "# See notebooks/03.preprocess_dataset.ipynb\nX_train_df, X_test_df, y_train, y_test = train_test_split(\n    X, y, test_size=0.20, random_state=42, stratify=y\n)\n# Fit imputation, IsolationForest, correlation pruning and StandardScaler on train only.\n# Oversample minority training classes to a maximum of 100,000 records per class.")
    elif text.startswith("import os\nimport time\nimport joblib"):
        set_text(p, "models = {\n    'Random Forest': RandomForestClassifier(n_estimators=200, random_state=42, n_jobs=-1),\n    'Logistic Regression': LogisticRegression(max_iter=3000, random_state=42),\n    'Decision Tree': DecisionTreeClassifier(random_state=42),\n}\n# Evaluate all candidates using weighted metrics, select the highest weighted F1,\n# and save it with joblib.dump(best_model, '../models/best_nids_model.joblib').")
    elif text.startswith("import React from \"react\""):
        set_text(p, "// Dashboard.tsx uploads a CSV through the FastAPI /predict endpoint.\n// Prediction results are shared through React DataContext and rendered by\n// StatCard, NetworkTrafficChart, AttackTypesChart, and ThreatAlertsTable.\n// Real-time capture uses the /realtime endpoints and the same result schema.")

# Replace mock dashboard ASCII with real screenshots.
dashboard_mock = next(p for p in doc.paragraphs if p.text.strip().startswith("AI-POWERED NIDS MONITORING HUB") or "AI-POWERED NIDS MONITORING HUB" in p.text)
replace_with_picture(dashboard_mock, ASSETS / "thesis_screenshots" / "dashboard.png")
alerts_mock = next(p for p in doc.paragraphs if "REAL-TIME THREAT ALERTS MANAGEMENT" in p.text)
replace_with_picture(alerts_mock, ASSETS / "thesis_screenshots" / "threat-alerts.png")

chapter6 = find_exact(doc, "CHAPTER 6: RESULTS AND EVALUATION")
insert_picture_before(chapter6, ASSETS / "thesis_screenshots" / "network-traffic.png", "Figure 5.5: Network Traffic Interface (Verified Application Capture)")
insert_picture_before(chapter6, ASSETS / "thesis_screenshots" / "settings.png", "Figure 5.6: Settings Interface (Verified Application Capture)")
insert_before(chapter6, "5.6 Implemented Security and Persistence Functions", "Heading 3")
insert_before(chapter6, "PostgreSQL persists operational data across sessions, including account information, network-flow records, predictions, alerts, block actions, and audit/history events. Authentication provides login and signup. Alert processing supports automatic policy-based blocking and explicit administrator-initiated manual blocking; each action is recorded for later review.", "Body Text")
insert_before(chapter6, "5.7 Testing", "Heading 3")
insert_before(chapter6, "The repository contains nine automated FastAPI tests covering health checks, model and artifact loading, valid prediction uploads, invalid extensions, schema validation, a SYN-flood signature rule, real-time status, CICFlowMeter aliases, and TShark interface parsing. The verified execution result was 9 passed with two dependency deprecation warnings in 2.82 seconds. No automated React component, end-to-end browser, or formal load-test suite was present in the inspected snapshot.", "Body Text")

# Chapter 6 verified results.
replace_start(doc, "Model performance was evaluated on a holdout", "Model performance was evaluated using an untouched stratified holdout partition containing 504,402 flow records and 47 selected features. Weighted precision, recall, and F1 are reported because the multiclass test distribution remains naturally imbalanced.")
replace_start(doc, "Four candidate algorithms", "Three candidate algorithms—Decision Tree, Random Forest, and Logistic Regression—were evaluated under the same preprocessing conditions. SVM, XGBoost, and CatBoost were not executed in the verified training notebook and are therefore excluded from the reported experimental results.")

fill_table(
    doc.tables[7],
    [
        ["Classifier Model", "Accuracy (%)", "Precision", "Recall", "F1-Score", "Training Time (s)", "Status"],
        ["Decision Tree", "99.7960", "0.997959", "0.997960", "0.997958", "71.04", "Selected/saved"],
        ["Random Forest", "99.7760", "0.997826", "0.997760", "0.997714", "399.88", "Evaluated"],
        ["Logistic Regression", "96.3474", "0.990236", "0.963474", "0.976133", "1198.67", "Evaluated"],
        ["SVM/XGBoost/CatBoost", "—", "—", "—", "—", "—", "Not run"],
    ],
)

cm_para = next(p for p in doc.paragraphs if p.text.strip().startswith("CONFUSION MATRIX"))
replace_with_picture(cm_para, ASSETS / "dissertation_assets" / "binary_confusion_matrix.png", width=5.8)
replace_start(doc, "Figure 6.1:", "Figure 6.1: Binary Confusion Matrix for the Final Saved Decision Tree Model")
replace_start(doc, "As detailed in Table 6.1", "Decision Tree produced the highest weighted F1-score and was saved as the deployed model. When its fifteen multiclass outputs were collapsed into BENIGN versus ATTACK, the confusion matrix contained TN=418,870, FP=357, FN=425, and TP=84,750. The corresponding binary accuracy was 0.998450, precision 0.995805, recall 0.995010, F1-score 0.995408, and ROC-AUC 0.997574.")

fill_table(
    doc.tables[8],
    [
        ["Threat Category", "Support", "Precision", "Recall", "F1-Score"],
        ["BENIGN", "419,227", "0.998986", "0.999148", "0.999067"],
        ["DDoS", "25,603", "0.999961", "0.999844", "0.999902"],
        ["DoS Hulk", "34,569", "0.997111", "0.998293", "0.997702"],
        ["PortScan", "18,164", "0.989052", "0.984805", "0.986924"],
        ["FTP-Patator", "1,187", "0.999158", "1.000000", "0.999579"],
        ["Web Attack – XSS", "130", "0.375887", "0.407692", "0.391144"],
    ],
)
replace_start(doc, "The system achieved detection rates exceeding", "Performance was strongest for DDoS, FTP-Patator, SSH-Patator, BENIGN, and high-volume DoS classes. Lower-frequency Web Attack classes were substantially harder to classify. Heartbleed had only two test samples and no remaining training samples after outlier filtering, producing zero precision, recall, and F1; this limitation must be considered when interpreting the high weighted average.")
replace_start(doc, "Inference latency was benchmarked", "A timed batch prediction of all 504,402 test flows using the saved Decision Tree completed in approximately 0.513 seconds on the verified machine. This corresponds to approximately 1.016 microseconds (0.001016 ms) per flow and about 983,922 flows per second. The measurement excludes CSV parsing, feature preprocessing, database writes, alert handling, blocking actions, and network/API overhead.")
replace_start(doc, "The empirical findings confirm", "The empirical findings show that all three tested models performed strongly on common classes, but Decision Tree achieved the best weighted F1 and the shortest training time. Weighted averages are influenced by the dominant BENIGN class, so per-class performance and the weak results for rare Web Attack and Heartbleed categories must also be reported.")
replace_start(doc, "Why Random Forest Outperformed", "Why Decision Tree Was Selected")
replace_start(doc, "While gradient boosting algorithms", "Decision Tree slightly exceeded Random Forest on accuracy and weighted F1 while training substantially faster (71.04 seconds versus 399.88 seconds). Logistic Regression required 1,198.67 seconds and produced lower accuracy. Because XGBoost, CatBoost, and SVM were not run, no comparative conclusion is made about those algorithms.")
replace_start(doc, "A detailed forensic examination", "Binary aggregation found 357 false positives and 425 false negatives. The multiclass report shows that the main limitations occur in rare categories, especially Web Attack – XSS, Web Attack – SQL Injection, Web Attack – Brute Force, Bot, and Heartbleed.")
replace_start(doc, "False Positive Drivers", "False Positive Interpretation: The available artifacts quantify false positives but do not preserve packet-level forensic explanations for each error. Causes such as unusual high-rate benign traffic should therefore be presented as hypotheses unless validated against the original flows.")
replace_start(doc, "False Negative Drivers", "False Negative Interpretation: The weaker recall for Web Attack classes is consistent with the limited number of test examples and the difficulty of distinguishing payload-oriented attacks using flow statistics alone. Deep packet inspection or application-layer features are proposed as future improvements.")

# Add verified model comparison figure before Chapter 7.
chapter7 = find_exact(doc, "CHAPTER 7: CONCLUSION AND RECOMMENDATIONS")
insert_picture_before(chapter7, ASSETS / "dissertation_assets" / "model_comparison.png", "Figure 6.2: Verified Accuracy and Weighted F1 Comparison")

# Conclusion and recommendations.
replace_start(doc, "By executing structured empirical surveys", "The implemented pipeline transformed the cleaned CICIDS2017 dataset into a leakage-controlled training matrix of 2,979,959 rows and an untouched test matrix of 504,402 rows, each with 47 selected features. Decision Tree was selected and serialized after achieving 0.997960 accuracy and 0.997958 weighted F1. The integrated system supports FastAPI prediction, live TShark/CICFlowMeter acquisition, PostgreSQL persistence, dashboard alerts, login/signup, automatic and manual blocking, and history/audit review.")
replace_start(doc, "For Network Administrators: Schedule monthly", "For Network Administrators: Review stored alert and block history regularly, validate automatic-block thresholds to reduce disruption from false positives, and retrain the model periodically through the notebook pipeline using approved and correctly labelled local data. A dashboard retraining endpoint should only be claimed after it is independently tested and deployed.")

# Technology table and preprocessing table.
fill_table(
    doc.tables[5],
    [
        ["Pipeline Stage", "Total Row Count", "Feature Count", "Data Condition"],
        ["Cleaned Dataset", "2,522,009", "70", "BENIGN plus 14 attack classes"],
        ["Initial Training Split", "2,017,607", "70", "80% stratified training partition"],
        ["Independent Test Split", "504,402", "70", "20% untouched holdout partition"],
        ["Train after Outlier Filter", "1,977,254", "70", "Isolation Forest applied to train only"],
        ["Correlation Pruned", "1,977,254", "47", "23 correlated features removed"],
        ["Oversampled Training Set", "2,979,959", "47", "Minority classes capped at 100,000"],
    ],
)
fill_table(
    doc.tables[6],
    [
        ["Layer / Category", "Technology / Library", "Purpose in Implementation"],
        ["Backend", "Python 3.14.3, FastAPI, Uvicorn", "REST API, validation, detection orchestration"],
        ["Machine Learning", "Scikit-learn, Joblib", "Training, preprocessing, evaluation, serialization"],
        ["Data Processing", "Pandas, NumPy, SciPy", "CSV and matrix transformations"],
        ["Persistence", "PostgreSQL", "Users, flows, predictions, alerts, blocks, and history"],
        ["Frontend", "React 18, TypeScript, Vite", "SPA dashboard and operational controls"],
        ["Capture/Visualization", "TShark, CICFlowMeter, Recharts", "Live capture, flow extraction, and charts"],
    ],
)

# Document metadata and save.
doc.core_properties.title = "AI-Based Network Intrusion Detection System with IoT Threat Monitoring Capabilities"
doc.core_properties.subject = "Revised dissertation aligned with verified SecureFlow-AI results and confirmed implemented features"
doc.core_properties.comments = "Model/data metrics verified from repository artifacts. PostgreSQL, alerts, blocking, authentication, and history recorded as user-confirmed implementation features."
doc.save(OUTPUT)
print(OUTPUT)
